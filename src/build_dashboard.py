"""Dựng dashboard Excel từ incidents_clean.csv (mọi số liệu trên Dashboard là công thức).

Chạy: python src/build_dashboard.py data/processed/incidents_clean.csv output/ServiceNow_Incident_SLA_Dashboard.xlsx
"""
import sys
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SRC, OUT = sys.argv[1], sys.argv[2]
inc = pd.read_csv(SRC, parse_dates=["Ngày mở", "Ngày giải quyết", "Ngày đóng"])
inc = inc.sort_values("Ngày mở")
N = len(inc)
last = N + 1

F, NAVY, LIGHT = "Arial", "1F3A5F", "E7F1FB"
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BODY = Font(name=F, size=10)
INPUT = Font(name=F, size=10, color="0000FF")
NOTE = Font(name=F, italic=True, size=9, color="595959")
RED = PatternFill("solid", fgColor="FBEBEB")
REDF = Font(name=F, color="C00000", bold=True)


def hdr(ws, row, col, labels):
    for i, t in enumerate(labels):
        c = ws.cell(row, col + i, t)
        c.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        c.fill, c.border = PatternFill("solid", fgColor=NAVY), BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def cell(ws, r, c, v, fmt=None, font=BODY):
    x = ws.cell(r, c, v)
    x.font, x.border = font, BOX
    if fmt:
        x.number_format = fmt
    return x


wb = Workbook()
db = wb.active
db.title = "Dashboard"
ds = wb.create_sheet("Data")

# ---------------- Data ----------------
raw_cols = list(inc.columns)  # A..M
calc_cols = ["Tháng mở", "Thời gian giải quyết (giờ)", "Nhóm chuyển", "Thiếu ngày giải quyết", "Thiếu nhóm xử lý", "Giải quyết ≤24h"]
hdr(ds, 1, 1, raw_cols + calc_cols)
date_fmt = "dd/mm/yyyy hh:mm"
for r, row in enumerate(inc.itertuples(index=False), start=2):
    for c, v in enumerate(row, start=1):
        if isinstance(v, float) and pd.isna(v) or v is pd.NaT:
            v = None
        elif isinstance(v, pd.Timestamp):
            v = v.to_pydatetime()
        x = ds.cell(r, c, v)
        if c in (2, 11, 12):
            x.number_format = date_fmt
    ds.cell(r, 14, f"=DATE(YEAR(B{r}),MONTH(B{r}),1)").number_format = "mm/yyyy"
    ds.cell(r, 15, f'=IF(K{r}="","",(K{r}-B{r})*24)').number_format = "0.0"
    ds.cell(r, 16, f'=IF(G{r}>=3,"3+",G{r}&"")')
    ds.cell(r, 17, f'=IF(K{r}="",1,0)')
    ds.cell(r, 18, f'=IF(E{r}="",1,0)')
    ds.cell(r, 19, f'=IF(O{r}="","",IF(O{r}<=24,1,0))')
for i, w in enumerate([13, 17, 14, 13, 12, 13, 10, 9, 11, 11, 17, 17, 8, 10, 12, 9, 11, 10, 10], start=1):
    ds.column_dimensions[ds.cell(1, i).column_letter].width = w
ds.freeze_panes = "B2"
ds.auto_filter.ref = f"A1:S{last}"

R = lambda col: f"Data!${col}$2:${col}${last}"  # noqa: E731

# ---------------- Dashboard ----------------
db.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGHIJKLM", (2, 24, 13, 13, 13, 13, 3, 16, 13, 13, 15, 15, 3)):
    db.column_dimensions[col].width = w
db["B1"] = "PHÂN TÍCH SLA XỬ LÝ SỰ CỐ IT – DỮ LIỆU SERVICENOW THỰC TẾ"
db["B1"].font = Font(name=F, bold=True, size=16, color=NAVY)
db["B2"] = (f"Nguồn: UCI ML Repository #498 (CC BY 4.0) – audit log ServiceNow của một công ty IT, "
            f"141.712 event → {N:,} incident, mở 02/2016–02/2017. Chi tiết xử lý: sheet Phuong_phap.").replace(",", ".")
db["B2"].font = NOTE
db["B3"] = "Mục tiêu SLA:"
db["B3"].font = Font(name=F, bold=True, size=10)
db["C3"] = 0.8
db["C3"].number_format = "0%"
db["C3"].font = INPUT
db["C3"].fill = PatternFill("solid", fgColor="FFFF00")
db["D3"] = "← giả định, sửa được; ô % đạt SLA dưới mục tiêu tô đỏ"
db["D3"].font = NOTE

sla = lambda extra="": f'COUNTIFS({R("J")},"Đạt"{extra})/COUNTIFS({R("A")},"<>"{extra})'  # noqa: E731
kpis = [
    ("Tổng incident", f'=COUNTA({R("A")})', "#,##0"),
    ("Tỷ lệ đạt SLA", f"=IFERROR({sla()},0)", "0.0%"),
    ("Giải quyết trong 24h", f'=IFERROR(SUM({R("S")})/COUNT({R("O")}),0)', "0.0%"),
    ("Có chuyển nhóm xử lý", f'=COUNTIF({R("G")},">=1")/COUNTA({R("A")})', "0.0%"),
    ("Từng chờ user bổ sung", f'=COUNTIF({R("I")},"Có")/COUNTA({R("A")})', "0.0%"),
    ("Thiếu ngày giải quyết", f'=SUM({R("Q")})', "#,##0"),
]
pos = [(5, 2), (5, 4), (5, 8), (8, 2), (8, 4), (8, 8)]
for (lab, f, fm), (r, c) in zip(kpis, pos):
    db.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + 1)
    db.merge_cells(start_row=r + 1, start_column=c, end_row=r + 1, end_column=c + 1)
    lc, vc = db.cell(r, c, lab), db.cell(r + 1, c, f)
    lc.font = Font(name=F, size=10, color="595959")
    vc.font, vc.number_format = Font(name=F, bold=True, size=20, color=NAVY), fm
    for x in (lc, vc):
        x.alignment = Alignment(horizontal="center", vertical="center")
    for rr in (r, r + 1):
        for cc in (c, c + 1):
            db.cell(rr, cc).fill = PatternFill("solid", fgColor=LIGHT)
    db.row_dimensions[r + 1].height = 30


def table(top, left, title, keys, key_col, metrics, key_fmt=None):
    db.cell(top, left, title).font = Font(name=F, bold=True, size=12, color=NAVY)
    hdr(db, top + 1, left, [m[0] for m in metrics])
    for i, k in enumerate(keys):
        r = top + 2 + i
        ref = f"${db.cell(r, left).column_letter}{r}"
        cell(db, r, left, k, key_fmt)
        for j, (_, fn, fm) in enumerate(metrics[1:], start=1):
            cell(db, r, left + j, fn(key_col, ref), fm)
    return top + 2, top + 1 + len(keys)


def crit(k, ref):
    return f',{R(k)},{ref}'


n_ = lambda k, ref: f"=COUNTIFS({R(k)},{ref})"  # noqa: E731
share = lambda k, ref: f'=COUNTIFS({R(k)},{ref})/COUNTA({R("A")})'  # noqa: E731
sla_ = lambda k, ref: f"=IFERROR({sla(crit(k, ref))},0)"  # noqa: E731
d24 = lambda k, ref: f'=IFERROR(SUMIFS({R("S")},{R(k)},{ref})/COUNTIFS({R(k)},{ref},{R("O")},">=0"),0)'  # noqa: E731
wait = lambda k, ref: f'=IFERROR(COUNTIFS({R(k)},{ref},{R("I")},"Có")/COUNTIFS({R(k)},{ref}),0)'  # noqa: E731
avg_re = lambda k, ref: f'=IFERROR(AVERAGEIF({R(k)},{ref},{R("G")}),0)'  # noqa: E731

std = [("", None, None), ("Số incident", n_, "#,##0"), ("% tổng", share, "0.0%"),
       ("% đạt SLA", sla_, "0.0%"), ("% giải quyết ≤24h", d24, "0.0%")]
t1a, t1b = table(12, 2, "1. Số lần chuyển nhóm xử lý", ["0", "1", "2", "3+"], "P",
                 [("Số lần chuyển",) + std[0][1:]] + std[1:])
t2a, t2b = table(12, 8, "2. Incident từng phải chờ user bổ sung thông tin", ["Không", "Có"], "I",
                 [("Từng chờ user",) + std[0][1:]] + std[1:])
prios = ["1 - Critical", "2 - High", "3 - Moderate", "4 - Low"]
t3a, t3b = table(19, 2, "3. Theo mức ưu tiên", prios, "C", [("Mức ưu tiên",) + std[0][1:]] + std[1:])
months = [datetime(2016, m, 1) for m in (3, 4, 5)]
t4a, t4b = table(17, 8, "4. Theo tháng mở (03–05/2016 chiếm 98%)", months, "N",
                 [("Tháng", None, None), ("Số incident", n_, "#,##0"), ("% tổng", share, "0.0%"),
                  ("% đạt SLA", sla_, "0.0%"), ("% giải quyết ≤24h", d24, "0.0%")], key_fmt="mm/yyyy")
top10 = inc["Nhóm xử lý"].value_counts().head(10).index.tolist()
t5a, t5b = table(26, 2, "5. Top 10 nhóm xử lý theo khối lượng (66% tổng incident)", top10, "E",
                 [("Nhóm xử lý", None, None), ("Số incident", n_, "#,##0"), ("% đạt SLA", sla_, "0.0%"),
                  ("% từng chờ user", wait, "0.0%"), ("TB lần chuyển nhóm", avg_re, "0.00")])

# 6. Data quality
dq_top = 23
db.cell(dq_top, 8, "6. Kiểm tra đầy đủ dữ liệu").font = Font(name=F, bold=True, size=12, color=NAVY)
hdr(db, dq_top + 1, 8, ["Trường", "Số dòng thiếu", "% thiếu"])
dq = [("Ngày giải quyết", "K"), ("Nhóm xử lý", "E"), ("Danh mục", "D")]
for i, (lab, col) in enumerate(dq):
    r = dq_top + 2 + i
    cell(db, r, 8, lab)
    cell(db, r, 9, f"=COUNTBLANK({R(col)})", "#,##0")
    cell(db, r, 10, f'=I{r}/COUNTA({R("A")})', "0.0%")

# SLA below target -> red
for col, a, b in (("E", t1a, t1b), ("K", t2a, t2b), ("E", t3a, t3b), ("K", t4a, t4b), ("D", t5a, t5b)):
    db.conditional_formatting.add(f"{col}{a}:{col}{b}",
                                  FormulaRule(formula=[f"{col}{a}<$C$3"], fill=RED, font=REDF))

# Key findings – text built from formulas so it stays in sync with the data
fr = t5b + 2
db.cell(fr, 2, "Nhận định chính").font = Font(name=F, bold=True, size=12, color=NAVY)
findings = [
    f'="• Incident không chuyển nhóm đạt SLA "&TEXT(E{t1a},"0%")&", chuyển ≥3 lần chỉ còn "&TEXT(E{t1b},"0%")&" → nên ưu tiên phân loại đúng nhóm ngay từ đầu."',
    f'="• Incident từng chờ user bổ sung thông tin đạt SLA "&TEXT(K{t2a+1},"0%")&" so với "&TEXT(K{t2a},"0%")&" → cần quy trình nhắc bổ sung thông tin sớm."',
    f'="• Critical/High đạt SLA rất thấp ("&TEXT(E{t3a},"0%")&" / "&TEXT(E{t3a+1},"0%")&") trong khi Low đạt "&TEXT(E{t3b},"0%")&" → cần xác minh lại cách đặt SLA cho nhóm ưu tiên cao."',
    f'="• "&TEXT(J{dq_top+2},"0.0%")&" incident đã đóng nhưng thiếu ngày giải quyết → không đo được thời gian xử lý, cần bổ sung khi đóng ticket."',
]
for i, f in enumerate(findings):
    db.cell(fr + 1 + i, 2, f).font = BODY
db.cell(fr + 5, 2, "Lưu ý: đây là tương quan quan sát được, chưa chứng minh nhân quả (ví dụ incident khó vừa bị chuyển nhóm nhiều vừa trễ SLA).").font = NOTE

# Charts
ch_row = fr + 7
c1 = BarChart()
c1.type, c1.title, c1.style = "col", "% đạt SLA theo số lần chuyển nhóm", 10
c1.add_data(Reference(db, min_col=5, min_row=t1a - 1, max_row=t1b), titles_from_data=True)
c1.set_categories(Reference(db, min_col=2, min_row=t1a, max_row=t1b))
c1.y_axis.numFmt, c1.y_axis.scaling.min, c1.y_axis.scaling.max = "0%", 0, 1
c1.x_axis.title = "Số lần chuyển nhóm"
c1.legend, c1.height, c1.width = None, 7.5, 13
db.add_chart(c1, f"B{ch_row}")
c2 = BarChart()
c2.type, c2.title, c2.style = "bar", "% đạt SLA – top 10 nhóm xử lý", 10
c2.add_data(Reference(db, min_col=4, min_row=t5a - 1, max_row=t5b), titles_from_data=True)
c2.set_categories(Reference(db, min_col=2, min_row=t5a, max_row=t5b))
c2.x_axis.numFmt = c2.y_axis.numFmt = "0%"
c2.y_axis.scaling.min, c2.y_axis.scaling.max = 0, 1
c2.x_axis.scaling.orientation = "maxMin"
c2.legend, c2.height, c2.width = None, 7.5, 14
db.add_chart(c2, f"H{ch_row}")
db.page_setup.orientation = "portrait"
db.page_setup.fitToWidth, db.page_setup.fitToHeight = 1, 1
db.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------- Phuong_phap ----------------
pp = wb.create_sheet("Phuong_phap")
rows = [
    ("Nguồn", "Incident management process enriched event log – UCI Machine Learning Repository #498, giấy phép CC BY 4.0. "
              "Audit log từ ServiceNow của một công ty IT, tên người/nhóm/danh mục đã được ẩn danh."),
    ("Link", "https://archive.ics.uci.edu/dataset/498/incident+management+process+enriched+event+log"),
    ("Tiền xử lý", "File gốc có 141.712 event (mỗi lần cập nhật ticket là 1 dòng). Gom thành 1 dòng/incident bằng Python pandas "
                   "(src/prepare_data.py): lấy event cuối cùng (sys_mod_count lớn nhất) làm trạng thái chốt; cờ 'Từng chờ user' = incident "
                   "có ít nhất 1 event ở trạng thái Awaiting User Info. Giá trị '?' trong file gốc coi là thiếu."),
    ("Cột tự tính", "N Tháng mở · O Thời gian giải quyết = (Ngày giải quyết − Ngày mở)×24 · P Nhóm chuyển 0/1/2/3+ · "
                    "Q, R cờ thiếu dữ liệu · S = 1 nếu giải quyết ≤24h."),
    ("% đạt SLA", "Số incident có made_sla = TRUE ở event cuối ÷ số incident trong nhóm. Ngưỡng SLA do hệ thống gốc đặt, dataset không công bố."),
    ("Hạn chế", "Không biết ngưỡng SLA cụ thể theo mức ưu tiên; 98% incident mở trong 03–05/2016, các tháng sau quá ít để so sánh; "
                "tháng 03/2016 có % đạt SLA thấp bất thường, chưa xác định nguyên nhân."),
]
hdr(pp, 1, 1, ["Mục", "Nội dung"])
for r, (a, b) in enumerate(rows, start=2):
    cell(pp, r, 1, a, font=Font(name=F, bold=True, size=10))
    cell(pp, r, 2, b).alignment = Alignment(wrap_text=True, vertical="top")
    pp.row_dimensions[r].height = 45
pp.column_dimensions["A"].width, pp.column_dimensions["B"].width = 14, 110

wb.save(OUT)
print("saved", OUT, N)
